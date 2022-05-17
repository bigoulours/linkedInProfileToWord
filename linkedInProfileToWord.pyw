from linkedin_api import Linkedin
import ttkbootstrap as ttk
from ttkbootstrap.scrolled import ScrolledText
from ttkbootstrap.tooltip import ToolTip
from tkinter import filedialog, messagebox
import os
from pathlib import Path
#import sys
import threading
import json
import urllib.request
import configparser
from docx.shared import Mm
from lib.linkedinToJSONresume import linkedin_to_json_resume
from docxtpl import DocxTemplate, InlineImage
from CI.version import SW_VERSION

config = configparser.ConfigParser()
config.read("linkedInProfileToWord.ini")
config_dict = {s:dict(config.items(s)) for s in config.sections()}

program_name = Path(__file__).stem + "-" + SW_VERSION

top = ttk.Window(themename=config_dict.get('General',{}).get('theme', 'cosmo'))
top.title(program_name)
top.geometry("800x680")
try:
    top.iconbitmap("images/linkedin.ico")
except:
    print("LinkedIn icon not found. Using default.")
status_str = ttk.StringVar(value="Ready")
template_path_str = ttk.StringVar(value="No template selected")
this_file_name = os.path.splitext(os.path.basename(__file__))[0]
resource_path = os.path.join(os.path.dirname(__file__), 'resources')
profile_to_export = dict()


def start_search(user, pwd, link):
    try:
        if not ("linkedin.com/in/" in link):
            status_str.set("Please enter a valid profile link!")
            return
        text_profile_json._text.configure(state="normal")
        text_profile_json.delete('1.0', 'end')
        text_profile_json._text.configure(state="disabled")
        export_to_word_btn.configure(state="disabled")
        profile_to_export.clear()
        public_id = link.split('linkedin.com/in/')[-1].strip('/')
        # Authenticate using any Linkedin account credentials
        try:
            api = Linkedin(user, pwd)
            status_str.set("Login successful!")
            top.update()
            export_to_word_btn.configure(state="disabled")
        except Exception as e:
            #exc_type, exc_obj, exc_tb = sys.exc_info()
            #fname = os.path.split(exc_tb.tb_frame.f_code.co_filename)[1]
            print("Error: " + repr(e)) #+ " in " + str(fname) + " line " + str(exc_tb.tb_lineno) + "\n")
            messagebox.showinfo("Error", "Login failed!\nCheck username and password.\n2FA must be disabled in LinkedIn settings.")
            return

        # see doc under https://linkedin-api.readthedocs.io/en/latest/api.html
        profile = api.get_profile(public_id=public_id)
        skills = api.get_profile_skills(public_id=public_id)
        contact_info = api.get_profile_contact_info(public_id=public_id)

        # #debug
        # print(json.dumps(profile, indent=4))
        # print(json.dumps(skills, indent=4))
        # print(json.dumps(contact_info, indent=4))

        profile_to_export.update(linkedin_to_json_resume(profile, skills, contact_info))
        #debug
        print(json.dumps(profile_to_export, indent=4))

        text_profile_json._text.configure(state="normal")
        text_profile_json.insert('insert', json.dumps(profile_to_export, indent=4))
        text_profile_json._text.configure(state="disabled")
        export_to_word_btn.configure(state="normal")
        status_str.set("Done")
        top.update()
    except Exception as e:
        #exc_type, exc_obj, exc_tb = sys.exc_info()
        #fname = os.path.split(exc_tb.tb_frame.f_code.co_filename)[1]
        print("Error: " + repr(e)) #+ " in " + str(fname) + " line " + str(exc_tb.tb_lineno) + "\n")
        status_str.set("Something went wrong! Check console output for more details")
        return


def create_start_search_thread(user, pwd, link):
    global search_thread
    if 'search_thread' in globals() and search_thread.is_alive():
        messagebox.showinfo("Search in progress", "Another search is still running.\nWait until it finishes or restart the program.")
        return
    search_thread = threading.Thread(target=start_search, args=(user, pwd, link))
    search_thread.daemon = True
    search_thread.start()


def load_template():
    chosen_file = filedialog.askopenfilename(filetypes=[("Word files", ".doc .docx")], initialdir=resource_path)
    if chosen_file is not None:
        template_path_str.set(chosen_file)
        return
    pass


def export_to_word(profile, template_path):
    status_str.set("Exporting to Word...")
    chosen_file = filedialog.asksaveasfile(mode='w', filetypes=[("Word files", ".docx")], defaultextension=".docx", initialfile=profile['basics']['name'])
    try:
        doc = DocxTemplate(template_path)
        if profile['basics']["image"]:
            urllib.request.urlretrieve(profile['basics']["image"], 'profile_pic_temp.jpg')
            image = InlineImage(doc, './profile_pic_temp.jpg', height=Mm(30))
        else:
            image = ""
        context = {'profile': profile, 'image': image}
        doc.render(context)
        doc.save(chosen_file.name)
        os.remove("profile_pic_temp.jpg")
        status_str.set("Word file saved under " + chosen_file.name)
    except Exception as e:
        #exc_type, exc_obj, exc_tb = sys.exc_info()
        #fname = os.path.split(exc_tb.tb_frame.f_code.co_filename)[1]
        print("Error: " + repr(e)) #+ " in " + str(fname) + " line " + str(exc_tb.tb_lineno) + "\n")
        status_str.set("Word file could not be saved!")


# Login frame
login_frame = ttk.Frame(top)
login_frame.pack(padx=10, pady=5, expand=False, fill="x")
label_usr = ttk.Label(login_frame, text="User")
label_usr.pack(side='left', expand=False)
ToolTip(label_usr, text="Your LinkedIn username\nPre-filled value can be changed in the .ini file")
entry_usr = ttk.Entry(login_frame)
last_login = config_dict.get('General',{}).get('user', '')
entry_usr.insert(0, last_login)
entry_usr.pack(side='left', expand=True, fill="x")
label_pwd = ttk.Label(login_frame, text="Pwd")
label_pwd.pack(side='left', expand=False)
entry_pwd = ttk.Entry(login_frame, show="*")
entry_pwd.pack(side='left', expand=True, fill="x")

separator = ttk.Separator(top, orient='horizontal')
separator.pack(side='top', pady=10, fill='x')

# Search filters 1
search_frame1 = ttk.Frame(top)
search_frame1.pack(padx=10, pady=0, side='top', fill="x")
label_profile_link = ttk.Label(search_frame1, text="Profile Link")
label_profile_link.pack(side='left', expand=False)
entry_profile_link = ttk.Entry(search_frame1)
entry_profile_link.pack(side='left', expand=True, fill="x")

# Buttons frame
btn_frame = ttk.Frame(top)
btn_frame.pack(padx=10, pady=0, side='top', fill="x")

start_search_btn = ttk.Button(btn_frame, text="Get JSON")
start_search_btn.pack(side='left', fill="none", expand=False)
start_search_btn['command'] = lambda: create_start_search_thread(entry_usr.get(), entry_pwd.get(), entry_profile_link.get())

btn_sub_frame = ttk.Frame(btn_frame)
btn_sub_frame.pack(side='left', fill="none", expand=True)
label_place_holder = ttk.Label(btn_sub_frame, text="")
label_place_holder.pack(side='top')
select_template_btn = ttk.Button(btn_sub_frame, text="Load template")
select_template_btn.pack(side='top', fill="none", expand=False)
select_template_btn['command'] = lambda: load_template()
label_template = ttk.Label(btn_sub_frame, textvariable=template_path_str)
label_template.pack(side='top', fill="x", expand=True)

export_to_word_btn = ttk.Button(btn_frame, text="Export to Word", state="disabled")
export_to_word_btn.pack(side='right', fill="none", expand=False)
export_to_word_btn['command'] = lambda: export_to_word(profile_to_export, template_path_str.get())

# Text Frame
text_frame = ttk.Frame(top)
text_frame.pack(padx=10, pady=10, side='top', fill="both", expand=True)
text_profile_json = ScrolledText(text_frame, font = ("Ubuntu", 12))
text_profile_json.pack(side='left', expand=True, fill="both")
text_profile_json._text.configure(state="disabled")

# Status frame
status_frame = ttk.Frame(top)
status_frame.pack(padx=10, pady=2, side='bottom', expand=False, fill="x")
label_status = ttk.Label(status_frame, textvariable=status_str)
label_status.pack(side='left', fill="x")

separator = ttk.Separator(top, orient='horizontal')
separator.pack(side='bottom', fill='x')

top.mainloop()


