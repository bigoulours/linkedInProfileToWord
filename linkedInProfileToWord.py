from linkedin_api import Linkedin
from tkinter import *
from tkinter import messagebox, ttk, filedialog, scrolledtext
import os
import threading
import json
import urllib.request
from docx.shared import Mm
from lib.linkedinToJSONresume import linkedin_to_json_resume
from docxtpl import DocxTemplate, InlineImage


top = Tk()
top.title("Linkedin Search")
window_width = 800
window_height = 680
top.geometry(str(window_width) + "x" + str(window_height))
try:
    top.iconbitmap("images/linkedin.ico")
except:
    print("LinkedIn icon not found. Using default.")
status_str = StringVar(value="Ready")
template_path_str = StringVar(value="No template selected")
this_file_name = os.path.splitext(os.path.basename(__file__))[0]
resource_path = os.path.join(os.path.dirname(__file__), 'resources')
profile_to_export = dict()


def start_search(user, pwd, link):
    try:
        if not ("linkedin.com/in/" in link):
            status_str.set("Please enter a valid profile link!")
            return
        text_profile_json.configure(state="normal")
        text_profile_json.delete('1.0', END)
        text_profile_json.configure(state="disabled")
        export_to_word_btn.configure(state="disabled")
        profile_to_export.clear()
        public_id = link.split('linkedin.com/in/')[-1].strip('/')
        # Authenticate using any Linkedin account credentials
        try:
            api = Linkedin(user, pwd)
            with open(this_file_name + '_login', 'w') as f:
                f.write(user)
            status_str.set("Login successful!")
            top.update()
            export_to_word_btn.configure(state="disabled")
        except Exception as e:
            exc_type, exc_obj, exc_tb = sys.exc_info()
            fname = os.path.split(exc_tb.tb_frame.f_code.co_filename)[1]
            print("Error: " + repr(e) + " in " + str(fname) + " line " + str(exc_tb.tb_lineno) + "\n")
            messagebox.showinfo("Error", "Login failed!\nCheck username and password.\n2FA must be disabled in LinkedIn settings.")
            return

        # see doc under https://linkedin-api.readthedocs.io/en/latest/api.html
        profile = api.get_profile(public_id=public_id)
        skills = api.get_profile_skills(public_id=public_id)
        contact_info = api.get_profile_contact_info(public_id=public_id)

        # #debug
        print(json.dumps(profile, indent=4))
        # print(json.dumps(skills, indent=4))
        # print(json.dumps(contact_info, indent=4))

        profile_to_export.update(linkedin_to_json_resume(profile, skills, contact_info))

        text_profile_json.configure(state="normal")
        text_profile_json.insert(INSERT, json.dumps(profile_to_export, indent=4))
        text_profile_json.configure(state="disabled")
        export_to_word_btn.configure(state="normal")
        status_str.set("Done")
        top.update()
    except Exception as e:
        exc_type, exc_obj, exc_tb = sys.exc_info()
        fname = os.path.split(exc_tb.tb_frame.f_code.co_filename)[1]
        print("Error: " + repr(e)+ " in " + str(fname) + " line " + str(exc_tb.tb_lineno) + "\n")
        status_str.set("Something went wrong! Check console output for more details")
        return


def create_start_search_thread(user, pwd, link):
    global search_thread
    if 'search_thread' in globals() and search_thread.is_alive():
        messagebox.showinfo("Search in progress", "Another search is still running.\nWait until it finishes or restart the program.")
        return
    search_thread = threading.Thread(target=start_search, args=(user, pwd, link))
    search_thread.daemon = True
    search_thread.start()


def load_template():
    chosen_file = filedialog.askopenfilename(filetypes=[("Word files", ".doc .docx")], initialdir=resource_path)
    if chosen_file is not None:
        template_path_str.set(chosen_file)
        return
    pass


def export_to_word(profile, template_path):
    status_str.set("Exporting to Word...")
    chosen_file = filedialog.asksaveasfile(mode='w', filetypes=[("Word files", ".docx")], defaultextension=".docx", initialfile=profile['basics']['name'])
    try:
        doc = DocxTemplate(template_path)
        if profile['basics']["image"]:
            urllib.request.urlretrieve(profile['basics']["image"], 'profile_pic_temp.jpg')
            image = InlineImage(doc, './profile_pic_temp.jpg', height=Mm(30))
        else:
            image = ""
        context = {'profile': profile, 'image': image}
        doc.render(context)
        doc.save(chosen_file.name)
        os.remove("profile_pic_temp.jpg")
        status_str.set("Word file saved under " + chosen_file.name)
    except Exception as e:
        exc_type, exc_obj, exc_tb = sys.exc_info()
        fname = os.path.split(exc_tb.tb_frame.f_code.co_filename)[1]
        print("Error: " + repr(e)+ " in " + str(fname) + " line " + str(exc_tb.tb_lineno) + "\n")
        status_str.set("Word file could not be saved!")


# Login frame
login_frame = Frame(top)
login_frame.pack(padx=10, pady=5, expand=False, fill="x")
label_usr = Label(login_frame, text="User")
label_usr.pack(side=LEFT, expand=False)
entry_usr = Entry(login_frame, bd=5)
last_login = ''
if os.path.isfile(this_file_name + '_login'):
    with open(this_file_name + '_login') as f:
        last_login = f.readlines()[0]
entry_usr.insert(0, last_login)
entry_usr.pack(side=LEFT, expand=True, fill="x")
label_pwd = Label(login_frame, text="Pwd")
label_pwd.pack(side=LEFT, expand=False)
entry_pwd = Entry(login_frame, show="*", bd=5)
entry_pwd.pack(side=LEFT, expand=True, fill="x")

separator = ttk.Separator(top, orient='horizontal')
separator.pack(side=TOP, pady=10, fill='x')

# Search filters 1
search_frame1 = Frame(top)
search_frame1.pack(padx=10, pady=0, side=TOP, fill="x")
label_profile_link = Label(search_frame1, text="Profile Link")
label_profile_link.pack(side=LEFT, expand=False)
entry_profile_link = Entry(search_frame1, bd=5)
entry_profile_link.pack(side=LEFT, expand=True, fill="x")

# Buttons frame
btn_frame = Frame(top)
btn_frame.pack(padx=10, pady=0, side=TOP, fill="x")

start_search_btn = Button(btn_frame, text="Get JSON")
start_search_btn.pack(side=LEFT, fill="none", expand=False)
start_search_btn['command'] = lambda: create_start_search_thread(entry_usr.get(), entry_pwd.get(), entry_profile_link.get())

btn_sub_frame = Frame(btn_frame)
btn_sub_frame.pack(side=LEFT, fill="none", expand=True)
label_place_holder = Label(btn_sub_frame, text="")
label_place_holder.pack(side=TOP)
select_template_btn = Button(btn_sub_frame, text="Load template")
select_template_btn.pack(side=TOP, fill="none", expand=False)
select_template_btn['command'] = lambda: load_template()
label_template = Label(btn_sub_frame, textvariable=template_path_str)
label_template.pack(side=TOP, fill="x", expand=True)

export_to_word_btn = Button(btn_frame, text="Export to Word", state="disabled")
export_to_word_btn.pack(side=RIGHT, fill="none", expand=False)
export_to_word_btn['command'] = lambda: export_to_word(profile_to_export, template_path_str.get())

# Text Frame
text_frame = Frame(top)
text_frame.pack(padx=10, pady=10, side=TOP, fill="both", expand=True)
text_profile_json = scrolledtext.ScrolledText(text_frame, bd=5)
text_profile_json.pack(side=LEFT, expand=True, fill="both")
text_profile_json.configure(state="disabled")

# Status frame
status_frame = Frame(top)
status_frame.pack(padx=10, pady=2, side=BOTTOM, expand=False, fill="x")
label_status = Label(status_frame, textvariable=status_str)
label_status.pack(side=LEFT, fill="x")

separator = ttk.Separator(top, orient='horizontal')
separator.pack(side=BOTTOM, fill='x')

top.mainloop()


